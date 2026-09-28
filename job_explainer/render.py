"""Render the job analysis and fourteen panels to a DOCX report + explanation.json.

Styling primitives (colors, fonts, headings, tables) live in
job_explainer/docx_style.py; this module is the document builder plus a
small Markdown-subset converter. Petal prompts are constrained (PRD §7 step
4) to exactly the Markdown this converter handles: paragraphs, `###`
headings, bulleted and numbered lists, bold, italic, and pipe tables —
nothing more, so there is nothing here that can silently drop content.
"""

from __future__ import annotations

import io
import json
import re
from datetime import date

import docx

from job_explainer import docx_style as style
from job_explainer.pipeline import PETALS, PetalResult

HEADING_PATTERN = re.compile(r"^#{1,3}\s+(.*)$")
BULLET_PATTERN = re.compile(r"^[-*]\s+(.*)$")
NUMBERED_PATTERN = re.compile(r"^\d+[.)]\s+(.*)$")
TABLE_ROW_PATTERN = re.compile(r"^\|(.+)\|$")
TABLE_SEPARATOR_PATTERN = re.compile(r"^\|[\s:|-]+\|$")
INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|_[^_]+?_)")


def slugify(text: str) -> str:
    text = (text or "").strip().lower()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-") or "untitled"


def report_filename(analysis: dict) -> str:
    slug = slugify(f"{analysis.get('title', '')}-{analysis.get('company', '')}")
    return f"job_explainer_{slug}.docx"


def _add_inline_runs(paragraph, text: str) -> None:
    for part in INLINE_PATTERN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif (part.startswith("*") and part.endswith("*")) or (part.startswith("_") and part.endswith("_")):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _strip_inline_markers(text: str) -> str:
    return INLINE_PATTERN.sub(lambda m: m.group(0).strip("*_"), text)


def _parse_table(lines: list[str]) -> tuple[list[str], list[list[str]]]:
    def split_row(line: str) -> list[str]:
        return [cell.strip() for cell in line.strip().strip("|").split("|")]

    header = split_row(lines[0])
    body_lines = lines[2:] if len(lines) > 1 and TABLE_SEPARATOR_PATTERN.match(lines[1]) else lines[1:]
    rows = [split_row(line) for line in body_lines]
    return header, rows


def markdown_to_docx(document: docx.Document, markdown_text: str) -> None:
    """Append the Markdown subset described above to `document`."""
    lines = (markdown_text or "").splitlines()
    i = 0
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        if paragraph_buffer:
            paragraph = document.add_paragraph()
            _add_inline_runs(paragraph, " ".join(paragraph_buffer))
            paragraph_buffer.clear()

    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip():
            flush_paragraph()
            i += 1
            continue

        heading_match = HEADING_PATTERN.match(line)
        if heading_match:
            flush_paragraph()
            style.add_heading(document, heading_match.group(1).strip(), level=3)
            i += 1
            continue

        if TABLE_ROW_PATTERN.match(line):
            flush_paragraph()
            table_lines = []
            while i < len(lines) and TABLE_ROW_PATTERN.match(lines[i].strip()):
                table_lines.append(lines[i].strip())
                i += 1
            header, rows = _parse_table(table_lines)
            style.add_table(document, header, [[_strip_inline_markers(c) for c in row] for row in rows])
            continue

        bullet_match = BULLET_PATTERN.match(line)
        if bullet_match:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline_runs(paragraph, bullet_match.group(1))
            i += 1
            continue

        numbered_match = NUMBERED_PATTERN.match(line)
        if numbered_match:
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Number")
            _add_inline_runs(paragraph, numbered_match.group(1))
            i += 1
            continue

        paragraph_buffer.append(line.strip())
        i += 1

    flush_paragraph()


def _add_summary_section(document: docx.Document, analysis: dict) -> None:
    style.add_heading(document, "Posting summary", level=1)

    facts = [
        ("Title", analysis.get("title")),
        ("Company", analysis.get("company")),
        ("Location", analysis.get("location")),
        ("Work arrangement", analysis.get("work_arrangement")),
        ("Seniority", analysis.get("seniority")),
        ("Sector", analysis.get("sector")),
    ]
    for label, value in facts:
        if not value:
            continue
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(str(value))

    if analysis.get("summary"):
        document.add_paragraph(analysis["summary"])

    if analysis.get("must_have_requirements"):
        style.add_heading(document, "Must-have requirements", level=2)
        for item in analysis["must_have_requirements"]:
            document.add_paragraph(str(item), style="List Bullet")

    if analysis.get("nice_to_have_requirements"):
        style.add_heading(document, "Nice-to-have requirements", level=2)
        for item in analysis["nice_to_have_requirements"]:
            document.add_paragraph(str(item), style="List Bullet")

    if analysis.get("tools_and_keywords"):
        paragraph = document.add_paragraph()
        paragraph.add_run("Tools and keywords: ").bold = True
        paragraph.add_run(", ".join(str(k) for k in analysis["tools_and_keywords"]))


def _add_panels(document: docx.Document, petal_results: dict[str, PetalResult]) -> None:
    for petal in PETALS:
        style.add_heading(document, f"{petal.number:02d}. {petal.title}", level=2)
        result = petal_results.get(petal.key)
        if result is None or result.error:
            paragraph = document.add_paragraph()
            message = result.error if result and result.error else "This panel wasn't generated."
            run = paragraph.add_run(message)
            run.italic = True
            run.font.color.rgb = style.MUTED_COLOR
            continue
        markdown_to_docx(document, result.markdown)


def render_report(analysis: dict, petal_results: dict[str, PetalResult]) -> docx.Document:
    """Build the full DOCX report: title block, posting summary, fourteen panels."""
    document = style.new_document()

    style.add_heading(document, analysis.get("title") or "Job explanation", level=0)
    today = date.today().strftime("%B %d, %Y").replace(" 0", " ")  # cross-platform: no %-d on Windows
    subtitle_parts = [p for p in [analysis.get("company"), today] if p]
    if subtitle_parts:
        paragraph = document.add_paragraph(" — ".join(subtitle_parts))
        paragraph.runs[0].italic = True
        paragraph.runs[0].font.color.rgb = style.MUTED_COLOR

    _add_summary_section(document, analysis)
    _add_panels(document, petal_results)

    return document


def build_docx_bytes(document: docx.Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def render_explanation_json(analysis: dict, petal_results: dict[str, PetalResult]) -> bytes:
    """A secondary, structured download: the analysis plus all panel Markdown."""
    panels = {}
    for petal in PETALS:
        result = petal_results.get(petal.key)
        panels[petal.key] = {
            "number": petal.number,
            "title": petal.title,
            "markdown": result.markdown if result else "",
            "error": result.error if result else "not generated",
        }
    payload = {"analysis": analysis, "panels": panels}
    return json.dumps(payload, indent=2).encode("utf-8")
